"""Read/write the applications Excel workbook.

Expected columns: role_title, company, job_link, resume_doc_link, status.
The bot fills in the `status` column as it works through the rows.
"""

from __future__ import annotations

from pathlib import Path

from applybot import config as C


class ExcelError(Exception):
    """Raised when the workbook is missing or malformed."""


def _require_openpyxl():
    try:
        import openpyxl  # noqa: F401
    except ImportError as exc:
        raise ExcelError(
            "openpyxl is not installed. Run: pip install -r requirements.txt"
        ) from exc


def load_applications(path: str | Path) -> tuple[list[dict], "WorkbookCtx"]:
    """Load rows from the workbook.

    Returns (rows, ctx) where each row is a dict with the column values plus
    `_excel_row` (the 1-based sheet row number) so status can be written back.
    """
    _require_openpyxl()
    import openpyxl

    path = Path(path)
    if not path.exists():
        raise ExcelError(f"Excel file not found: {path}")

    wb = openpyxl.load_workbook(path)
    ws = wb.active

    headers = [c.value for c in ws[1]]
    missing = [c for c in C.REQUIRED_COLUMNS if c not in headers]
    if missing:
        raise ExcelError(
            f"{path} is missing required columns: {missing}. "
            f"Found: {[h for h in headers if h]}"
        )

    # Make sure a status column exists (append it if absent).
    if C.COL_STATUS not in headers:
        ws.cell(row=1, column=len(headers) + 1, value=C.COL_STATUS)
        headers.append(C.COL_STATUS)

    col_idx = {name: headers.index(name) for name in headers if name}
    rows: list[dict] = []
    for r in range(2, ws.max_row + 1):
        role = ws.cell(row=r, column=col_idx[C.COL_ROLE] + 1).value
        company = ws.cell(row=r, column=col_idx[C.COL_COMPANY] + 1).value
        if not role and not company:
            continue  # skip blank rows
        row = {"_excel_row": r}
        for name, i in col_idx.items():
            row[name] = ws.cell(row=r, column=i + 1).value
        rows.append(row)

    return rows, _WorkbookCtx(wb, ws, col_idx)


class _WorkbookCtx:
    """Holds the open workbook so the runner can write statuses back."""

    def __init__(self, wb, ws, col_idx):
        self.wb = wb
        self.ws = ws
        self.col_idx = col_idx
        self.path: Path | None = None

    def set_status(self, excel_row: int, status: str) -> None:
        col = self.col_idx[C.COL_STATUS] + 1
        self.ws.cell(row=excel_row, column=col, value=status)

    def save(self, path: str | Path) -> None:
        self.wb.save(path)


def create_sample_excel(path: str | Path) -> Path:
    """Write a sample applications.xlsx with two example rows."""
    _require_openpyxl()
    import openpyxl

    path = Path(path)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "applications"
    ws.append(
        [C.COL_ROLE, C.COL_COMPANY, C.COL_JOB_LINK, C.COL_RESUME_LINK, C.COL_STATUS]
    )
    ws.append(
        [
            "Data Scientist",
            "ExampleCorp",
            "https://example.com/careers/data-scientist-123",
            "https://docs.google.com/document/d/PASTE_DOC_ID_HERE/edit",
            C.STATUS_PENDING,
        ]
    )
    ws.append(
        [
            "Machine Learning Engineer",
            "SampleAI",
            "https://example.com/jobs/ml-engineer-456",
            "https://docs.google.com/document/d/PASTE_DOC_ID_HERE_2/edit",
            C.STATUS_PENDING,
        ]
    )
    # Sensible column widths for readability.
    for col, width in zip("ABCDE", [28, 18, 50, 55, 16]):
        ws.column_dimensions[col].width = width
    wb.save(path)
    return path
